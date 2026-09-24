"""Build Sequence #TBD -- Phase 6, evidence file text extraction. Extracts
representative text content from a project evidence file's bytes, dispatched
by file extension. This is Phase A of the planned two-phase build (Ankita,
22-23 Sept 2026): PDF, DOCX, PPTX, XLSX, CSV, TXT supported here; images
(needing OCR) and email formats (.msg/.eml) are explicitly deferred to a
later Phase B, not attempted by this module.

Deliberately extracts REPRESENTATIVE content, not necessarily the full
document -- for PDF/DOCX/PPTX, a bounded number of pages/slides is read
(enough to understand what the document is and roughly what it covers, per
Ankita's own scoping: "what the document is, its type... actual content, so
that we can simply validate the relevance and map it"), not the complete
text of a potentially very large file. For XLSX/CSV, header rows plus a
bounded sample of data rows serve the same purpose.
"""
import io

MAX_PDF_PAGES = 10
MAX_PPTX_SLIDES = 15
MAX_DOCX_PARAGRAPHS = 200
MAX_XLSX_ROWS_PER_SHEET = 20
MAX_CSV_ROWS = 20


def extract_representative_text(file_bytes: bytes, filename: str) -> dict:
    """Returns {"status": "success", "file_type": ..., "text": ...} on success,
    or {"status": "unsupported", "file_type": ..., "message": ...} for a file
    type Phase A does not yet handle (images, email formats, video, etc.) --
    never raises for an unsupported type, since a file this can't read should
    still exist as a ProjectEvidenceFile row with a clear status, not crash
    the upload/mapping flow around it.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    try:
        if ext == "pdf":
            return {"status": "success", "file_type": "pdf", "text": _extract_pdf(file_bytes)}
        elif ext in ("docx",):
            return {"status": "success", "file_type": "docx", "text": _extract_docx(file_bytes)}
        elif ext in ("pptx",):
            return {"status": "success", "file_type": "pptx", "text": _extract_pptx(file_bytes)}
        elif ext in ("xlsx",):
            return {"status": "success", "file_type": "xlsx", "text": _extract_xlsx(file_bytes)}
        elif ext == "csv":
            return {"status": "success", "file_type": "csv", "text": _extract_csv(file_bytes)}
        elif ext == "txt":
            return {"status": "success", "file_type": "txt", "text": _extract_txt(file_bytes)}
        else:
            return {
                "status": "unsupported",
                "file_type": ext or "unknown",
                "message": f"File type '.{ext}' is not yet supported for content extraction "
                           f"(Phase A covers pdf, docx, pptx, xlsx, csv, txt only).",
            }
    except Exception as e:
        return {
            "status": "error",
            "file_type": ext or "unknown",
            "message": f"Failed to extract content: {str(e)}",
        }


def _extract_pdf(file_bytes: bytes) -> str:
    import fitz
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_to_read = min(MAX_PDF_PAGES, len(doc))
    text_parts = []
    for i in range(pages_to_read):
        page_text = doc[i].get_text()
        if page_text.strip():
            text_parts.append(f"[Page {i + 1}]\n{page_text.strip()}")
    doc.close()
    return "\n\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    for i, para in enumerate(doc.paragraphs):
        if i >= MAX_DOCX_PARAGRAPHS:
            break
        if para.text.strip():
            paragraphs.append(para.text.strip())
    return "\n".join(paragraphs)


def _extract_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    slide_texts = []
    for i, slide in enumerate(prs.slides):
        if i >= MAX_PPTX_SLIDES:
            break
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            slide_texts.append(f"[Slide {i + 1}]\n" + "\n".join(texts))
    return "\n\n".join(slide_texts)


def _extract_xlsx(file_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet_texts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_XLSX_ROWS_PER_SHEET:
                break
            row_str = ", ".join(str(c) for c in row if c is not None)
            if row_str.strip():
                rows_text.append(row_str)
        if rows_text:
            sheet_texts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows_text))
    wb.close()
    return "\n\n".join(sheet_texts)


def _extract_csv(file_bytes: bytes) -> str:
    import csv
    text = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for i, row in enumerate(reader):
        if i >= MAX_CSV_ROWS:
            break
        rows.append(", ".join(row))
    return "\n".join(rows)


def _extract_txt(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    # Bound plain text the same way as other types -- a very large log file
    # shouldn't be read in full for a relevance judgment.
    return text[:20000]
